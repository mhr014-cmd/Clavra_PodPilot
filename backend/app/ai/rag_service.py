"""
RAG Service — Clavra ProdPilot™
Embed, store, and retrieve knowledge from uploaded documents using pgvector.
Supports PDF, TXT, DOCX, XLSX.
Falls back to keyword search when OpenAI embeddings are unavailable.
"""
import os, json
from typing import Optional
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import text, select
from app.config import settings

EMBEDDING_MODEL = "text-embedding-3-small"
EMBEDDING_DIM   = 1536
CHUNK_SIZE      = 512    # words per chunk
CHUNK_OVERLAP   = 64
TOP_K           = 5
MIN_SIMILARITY  = 0.72

RAG_ANSWER_PROMPT = """You are a knowledgeable assistant for a manufacturing company.
Answer the question using ONLY the provided document excerpts.
Always cite your sources with [Document: name, Page: X].
If the answer is not in the documents, say so clearly.
Be concise and practical."""


# ── Embedding ─────────────────────────────────────────────────────────────

async def embed_text(text_content: str) -> list[float]:
    """Embed text using OpenAI text-embedding-3-small (1536 dims).
    Returns zero vector if OpenAI is unavailable — so ingest always completes
    and keyword search still works on the stored chunks."""
    if settings.OPENAI_API_KEY and settings.OPENAI_API_KEY.startswith("sk-"):
        try:
            from openai import AsyncOpenAI
            client = AsyncOpenAI(api_key=settings.OPENAI_API_KEY)
            response = await client.embeddings.create(
                model=EMBEDDING_MODEL,
                input=text_content[:8000],
            )
            return response.data[0].embedding
        except Exception:
            pass
    return [0.0] * EMBEDDING_DIM


# ── Shared INSERT helper ──────────────────────────────────────────────────

async def _insert_chunk(db: AsyncSession, doc_id: int, org_id: int,
                        chunk_idx: int, page_num: int, chunk: str) -> None:
    embedding = await embed_text(chunk)
    embedding_str = f"[{','.join(map(str, embedding))}]"
    # Use CAST(...AS vector) instead of ::vector — SQLAlchemy text() misparses ::
    await db.execute(text("""
        INSERT INTO document_chunks
            (document_id, org_id, chunk_index, page_number, content, embedding, token_count)
        VALUES
            (:doc_id, :org_id, :chunk_index, :page_number, :content,
             CAST(:embedding AS vector), :token_count)
    """), {
        "doc_id": doc_id, "org_id": org_id,
        "chunk_index": chunk_idx, "page_number": page_num,
        "content": chunk,
        "embedding": embedding_str,
        "token_count": len(chunk.split()),
    })


# ── PDF ingestion ─────────────────────────────────────────────────────────

async def ingest_pdf(file_path: str, doc_id: int, org_id: int, db: AsyncSession) -> int:
    """Extract text from PDF, chunk, embed, store in DB."""
    import fitz  # PyMuPDF
    doc = fitz.open(file_path)
    total_chunks = 0
    for page_num, page in enumerate(doc, start=1):
        page_text = page.get_text().strip()
        if not page_text:
            continue
        for chunk_idx, chunk in enumerate(_chunk_text(page_text, CHUNK_SIZE)):
            if not chunk.strip():
                continue
            await _insert_chunk(db, doc_id, org_id, chunk_idx, page_num, chunk)
            total_chunks += 1
    await db.commit()
    doc.close()
    return total_chunks


# ── TXT ingestion ─────────────────────────────────────────────────────────

async def ingest_txt(file_path: str, doc_id: int, org_id: int, db: AsyncSession) -> int:
    """Extract text from plain .txt file, chunk, embed, store."""
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        full_text = f.read().strip()
    if not full_text:
        return 0
    total_chunks = 0
    for chunk_idx, chunk in enumerate(_chunk_text(full_text, CHUNK_SIZE)):
        if not chunk.strip():
            continue
        await _insert_chunk(db, doc_id, org_id, chunk_idx, 1, chunk)
        total_chunks += 1
    await db.commit()
    return total_chunks


# ── DOCX ingestion ────────────────────────────────────────────────────────

async def ingest_docx(file_path: str, doc_id: int, org_id: int, db: AsyncSession) -> int:
    """Extract text from .docx using python-docx, chunk, embed, store."""
    from docx import Document as DocxDocument
    doc = DocxDocument(file_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(paragraphs)
    if not full_text:
        return 0
    total_chunks = 0
    for chunk_idx, chunk in enumerate(_chunk_text(full_text, CHUNK_SIZE)):
        if not chunk.strip():
            continue
        await _insert_chunk(db, doc_id, org_id, chunk_idx, 1, chunk)
        total_chunks += 1
    await db.commit()
    return total_chunks


# ── XLSX ingestion ────────────────────────────────────────────────────────

async def ingest_xlsx(file_path: str, doc_id: int, org_id: int, db: AsyncSession) -> int:
    """Extract text from Excel .xlsx using openpyxl, chunk, embed, store."""
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    text_parts: list[str] = []
    for sheet in wb.worksheets:
        text_parts.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_str = "\t".join(str(c) if c is not None else "" for c in row)
            if row_str.strip():
                text_parts.append(row_str)
    wb.close()
    full_text = "\n".join(text_parts)
    if not full_text.strip():
        return 0
    total_chunks = 0
    for chunk_idx, chunk in enumerate(_chunk_text(full_text, CHUNK_SIZE)):
        if not chunk.strip():
            continue
        await _insert_chunk(db, doc_id, org_id, chunk_idx, 1, chunk)
        total_chunks += 1
    await db.commit()
    return total_chunks


# ── Chunker ───────────────────────────────────────────────────────────────

def _chunk_text(text_content: str, chunk_size: int) -> list[str]:
    """Split text into overlapping chunks by word count."""
    words = text_content.split()
    chunks = []
    step = chunk_size - CHUNK_OVERLAP
    for i in range(0, len(words), step):
        chunk = " ".join(words[i:i + chunk_size])
        if chunk:
            chunks.append(chunk)
        if i + chunk_size >= len(words):
            break
    return chunks


# ── Keyword fallback search ───────────────────────────────────────────────

async def _keyword_search(question: str, org_id: int, db: AsyncSession) -> list:
    """PostgreSQL ILIKE keyword search — used when embedding is unavailable."""
    words = [w.strip() for w in question.split() if len(w.strip()) >= 3][:8]
    if not words:
        return []
    conditions = " OR ".join(f"dc.content ILIKE :kw{i}" for i in range(len(words)))
    params: dict = {"org_id": org_id, "top_k": TOP_K}
    for i, w in enumerate(words):
        params[f"kw{i}"] = f"%{w}%"
    rows = await db.execute(text(f"""
        SELECT dc.content, dc.page_number, kd.original_name, kd.doc_type,
               0.5 AS similarity
        FROM document_chunks dc
        JOIN knowledge_documents kd ON dc.document_id = kd.id
        WHERE dc.org_id = :org_id AND ({conditions})
        ORDER BY dc.id
        LIMIT :top_k
    """), params)
    return rows.fetchall()


# ── Vector search ─────────────────────────────────────────────────────────

async def query_knowledge_base(question: str, org_id: int, db: Optional[AsyncSession] = None) -> dict:
    """
    Full RAG pipeline: embed question → vector search → generate answer with citations.
    Falls back to keyword search if embeddings are unavailable.
    """
    from app.database import AsyncSessionLocal

    close_db = False
    if db is None:
        db = AsyncSessionLocal()
        close_db = True

    using_keyword_fallback = False

    try:
        # Step 1: Try vector search
        chunks = []
        try:
            q_embedding = await embed_text(question)
            # Only use vector search if we have a real embedding (not all zeros)
            if any(v != 0.0 for v in q_embedding[:10]):
                embedding_str = f"[{','.join(map(str, q_embedding))}]"
                rows = await db.execute(text("""
                    SELECT
                        dc.content,
                        dc.page_number,
                        kd.original_name,
                        kd.doc_type,
                        1 - (dc.embedding <=> CAST(:q_vec AS vector)) AS similarity
                    FROM document_chunks dc
                    JOIN knowledge_documents kd ON dc.document_id = kd.id
                    WHERE dc.org_id = :org_id
                      AND 1 - (dc.embedding <=> CAST(:q_vec AS vector)) > :min_sim
                    ORDER BY similarity DESC
                    LIMIT :top_k
                """), {
                    "q_vec": embedding_str,
                    "org_id": org_id,
                    "min_sim": MIN_SIMILARITY,
                    "top_k": TOP_K,
                })
                chunks = rows.fetchall()
        except Exception:
            pass

        # Step 2: Fall back to keyword search if vector search returned nothing
        if not chunks:
            chunks = await _keyword_search(question, org_id, db)
            using_keyword_fallback = bool(chunks)

        if not chunks:
            return {
                "answer": (
                    "I couldn't find relevant information in the uploaded documents.\n\n"
                    "Please ensure the relevant document has been uploaded to the Knowledge Base."
                ),
                "sources": [],
            }

        # Step 3: Build context
        context_parts = []
        sources = []
        for row in chunks:
            context_parts.append(
                f"[Document: {row.original_name}, Page {row.page_number}]\n{row.content}"
            )
            sources.append({
                "document":   row.original_name,
                "page":       row.page_number,
                "similarity": round(float(row.similarity), 3),
                "doc_type":   row.doc_type,
            })
        context = "\n\n---\n\n".join(context_parts)

        # Step 4: Generate answer
        answer = await _generate_rag_answer(question, context)
        if using_keyword_fallback:
            answer += "\n\n*Note: Semantic search unavailable — results based on keyword matching.*"

        return {"answer": answer, "sources": sources}

    finally:
        if close_db:
            await db.close()


async def _generate_rag_answer(question: str, context: str) -> str:
    """Generate a cited answer from retrieved context.
    Tries OpenAI GPT-4o → Ollama llm → formatted excerpt fallback.
    """
    prompt = f"Context from documents:\n\n{context}\n\nQuestion: {question}"

    # ── OpenAI ────────────────────────────────────────────────────────────
    if settings.OPENAI_API_KEY and settings.OPENAI_API_KEY.startswith("sk-"):
        try:
            from openai import AsyncOpenAI
            client = AsyncOpenAI(api_key=settings.OPENAI_API_KEY)
            response = await client.chat.completions.create(
                model="gpt-4o",
                messages=[
                    {"role": "system", "content": RAG_ANSWER_PROMPT},
                    {"role": "user",   "content": prompt}
                ],
                temperature=0.2,
                max_tokens=600,
                timeout=10.0,
            )
            return response.choices[0].message.content.strip()
        except Exception:
            pass

    # ── Ollama fallback ───────────────────────────────────────────────────
    try:
        from app.ai.ollama_service import ask_ollama
        ollama_prompt = (
            f"{RAG_ANSWER_PROMPT}\n\n"
            f"Context from documents:\n{context[:2000]}\n\n"
            f"Question: {question}\n\n"
            f"Answer concisely using only the document content above:"
        )
        answer = await ask_ollama(ollama_prompt)
        if answer and len(answer.strip()) > 30:
            return answer.strip() + "\n\n*Answer generated by local AI model.*"
    except Exception:
        pass

    # ── Plain excerpt fallback ─────────────────────────────────────────────
    lines = []
    for part in context.split("\n\n---\n\n")[:3]:
        header_end = part.find("\n")
        if header_end > 0:
            header = part[:header_end].strip()
            body   = part[header_end:].strip()[:400]
            lines.append(f"**{header}**\n{body}")
        else:
            lines.append(part[:400])
    return "**Relevant document content:**\n\n" + "\n\n".join(lines)
